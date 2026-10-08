# Genera el plan de marketing completo en Excel a partir de src/plan-marketing.js y src/top8020.js
# Uso: python3 src/generar-plan-marketing.py Plan_Marketing_CLT_Mexico_Latam.xlsx
import json, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

def node(mod):
    return json.loads(subprocess.check_output(["node", "-e", f"console.log(JSON.stringify(require('./src/{mod}')))"]))

plan = node("plan-marketing")["es"]
top = node("top8020")
body = plan["body"]

INK, PINK, LIGHT, GRID = "1B1F3B", "FCE4EC", "EEF2F7", "C9D1DC"
HEAD = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor=INK)
TITLE = Font(bold=True, size=16, color=INK); SUB = Font(italic=True, size=10, color="555555")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color=GRID); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT = PatternFill("solid", fgColor="FFF8E1")  # celdas para capturar

def section(h1):
    """Bloques entre un h1 y el siguiente."""
    i = next(k for k, b in enumerate(body) if b.get("h1", "").startswith(h1))
    out = []
    for b in body[i + 1:]:
        if "h1" in b: break
        out.append(b)
    return out

def sheet(wb, name, title, subtitle=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TITLE
    if subtitle: ws["A2"] = subtitle; ws["A2"].font = SUB
    ws.sheet_view.showGridLines = False
    return ws

def table(ws, row, cols, rows, widths=None, total=False):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c); cell.font = HEAD; cell.fill = HFILL; cell.alignment = WRAP; cell.border = BOX
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r, 1):
            cell = ws.cell(row + i, j, v); cell.alignment = WRAP; cell.border = BOX
            if j == 1: cell.font = Font(bold=True)
            if i % 2 == 0: cell.fill = PatternFill("solid", fgColor=LIGHT)
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
    return row + len(rows) + 1

def blocks(ws, row, bl, widths):
    for b in bl:
        if "h2" in b:
            ws.cell(row, 1, b["h2"]).font = Font(bold=True, size=12, color="C2185B"); row += 1
        elif "p" in b or "note" in b:
            ws.cell(row, 1, b.get("p") or b.get("note")).alignment = WRAP
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(4, len(widths)))
            ws.row_dimensions[row].height = 45 if "p" in b else 30
            if "note" in b: ws.cell(row, 1).font = SUB
            row += 2
        elif "bullets" in b:
            for t in b["bullets"]:
                ws.cell(row, 1, "• " + t).alignment = WRAP
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(4, len(widths)))
                ws.row_dimensions[row].height = 32; row += 1
            row += 1
        elif "table" in b:
            t = b["table"]
            row = table(ws, row, t["cols"], t["rows"], widths) + 1
    for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
    return row

wb = Workbook(); wb.remove(wb.active)

# 1. Resumen
ws = sheet(wb, "Resumen", plan["meta"]["title"], plan["meta"]["subtitle"])
ws["A3"] = plan["meta"]["preparedBy"] + " · " + plan["meta"]["date"]; ws["A3"].font = SUB
r = blocks(ws, 5, section("1."), [120])
ws.cell(r, 1, "Contenido del libro").font = Font(bold=True, size=12, color="C2185B"); r += 1
TABS = ["Punto de partida", "Objetivos", "Segmentos", "Compradores", "Posicionamiento", "Mezcla 4P", "Fases México",
        "ABM 17 cuentas", "Canales", "Contenido", "Eventos", "Embudo", "Presupuesto", "KPIs mensuales",
        "Calendario", "Latinoamérica", "Pedir a CLT", "Riesgos"]
for t in TABS:
    c = ws.cell(r, 1, t); c.hyperlink = f"#'{t}'!A1"; c.font = Font(color="0079A8", underline="single"); r += 1
ws.cell(r + 1, 1, plan["meta"]["estimateNote"]).font = SUB

# 2–7: secciones de texto y tablas
for name, h1, title, widths in [
    ("Punto de partida", "2.", "Punto de partida", [22, 60, 55]),
    ("Objetivos", "3.", "Objetivos (metas por año)", [48, 16, 22, 18]),
    ("Segmentos", "4.", "Mercado meta y segmentos", [24, 45, 38, 40, 14]),
    ("Compradores", "5.", "Compradores: quién decide y qué le decimos", [26, 34, 42, 48]),
    ("Posicionamiento", "6.", "Posicionamiento y marca", [22, 45, 60]),
    ("Mezcla 4P", "7.", "Mezcla de marketing (4P)", [16, 120]),
    ("Fases México", "8.", "Plan por fases en México", [18, 10, 30, 80]),
]:
    ws = sheet(wb, name, title)
    blocks(ws, 3, section(h1), widths)

# 8. ABM: rastreador de las 17 cuentas
ws = sheet(wb, "ABM 17 cuentas", "Marketing por cuentas (ABM) — 17 cuentas 80/20", "Pasos: 1 Investigar · 2 Muestra personalizada · 3 Visita/demo · 4 Prueba piloto · 5 Propuesta. Capture fechas en las celdas amarillas.")
r = blocks(ws, 4, section("9."), [6, 34, 18, 16, 40, 14, 14, 14, 14, 14, 16, 30])
cols = ["#", "Cuenta", "Segmento", "Ciudad", "Por qué es prioridad", "1 Investigación", "2 Muestra enviada", "3 Demo", "4 Piloto", "5 Propuesta", "Estatus", "Próximo paso / responsable"]
rows = [[i + 1, a[0], "Prospecto", a[1], a[2]] + [""] * 5 + ["Por contactar", ""] for i, a in enumerate(top["prospects"])]
rows += [[len(rows) + i + 1, a[0], "OEM", a[1], a[2] + " — " + a[4]] + [""] * 5 + ["Por contactar", ""] for i, a in enumerate(top["oems"])]
start = r + 1
end = table(ws, start, cols, rows) - 1
for rr in range(start + 1, end + 1):
    for cc in range(6, 11): ws.cell(rr, cc).fill = INPUT; ws.cell(rr, cc).number_format = "dd/mm/yyyy"
    ws.cell(rr, 11).fill = INPUT; ws.cell(rr, 12).fill = INPUT
dv = DataValidation(type="list", formula1='"Por contactar,Muestra enviada,Demo agendada,Piloto,Propuesta,Ganado,Perdido,En pausa"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"K{start + 1}:K{end}")
ws.cell(end + 2, 2, "Cuentas con demo o piloto").font = Font(bold=True)
ws.cell(end + 2, 5, f'=COUNTA(H{start + 1}:H{end})+COUNTA(I{start + 1}:I{end})-SUMPRODUCT((H{start + 1}:H{end}<>"")*(I{start + 1}:I{end}<>""))')
ws.cell(end + 3, 2, "Meta año 1").font = Font(bold=True); ws.cell(end + 3, 5, 10)
ws.freeze_panes = ws.cell(start + 1, 3)

# 9–11
for name, h1, title, widths in [
    ("Canales", "10.", "Canales y tácticas", [26, 28, 60, 32]),
    ("Contenido", "11.", "Contenido", [28, 60, 26]),
    ("Eventos", "12.", "Eventos y ferias", [30, 34, 46, 12]),
]:
    ws = sheet(wb, name, title)
    blocks(ws, 3, section(h1), widths)

# 12. Embudo con fórmulas
ws = sheet(wb, "Embudo", "Embudo comercial año 1", "Las cifras amarillas son supuestos editables; las conversiones se recalculan solas.")
cols = ["Etapa", "Meta año 1", "Conversión vs. etapa anterior", "Real", "Avance", "Responsable"]
stages = [("Contactos generados (ferias, digital, ABM)", 900, "Marketing"), ("Leads calificados (MQL)", 300, "Marketing + atención al cliente"),
          ("Demostraciones o pruebas", 40, "Vendedor + técnico"), ("Propuestas formales", 12, "Vendedor"), ("Líneas híbridas vendidas", 2, "Dirección + vendedor")]
table(ws, 4, cols, [[s, n, "", "", "", o] for s, n, o in stages], [44, 14, 26, 12, 12, 32])
for i in range(len(stages)):
    rr = 5 + i
    ws.cell(rr, 2).fill = INPUT; ws.cell(rr, 4).fill = INPUT
    if i: ws.cell(rr, 3, f"=IFERROR(B{rr}/B{rr - 1},0)").number_format = "0%"
    ws.cell(rr, 5, f'=IFERROR(D{rr}/B{rr},"")').number_format = "0%"
rr = 5 + len(stages) + 1
ws.cell(rr, 1, "Talleres con consumibles").font = Font(bold=True); ws.cell(rr, 2, 30).fill = INPUT
ws.cell(rr, 3, f"=IFERROR(B{rr}/B6,0)").number_format = "0%"; ws.cell(rr, 4).fill = INPUT
ws.cell(rr, 5, f'=IFERROR(D{rr}/B{rr},"")').number_format = "0%"; ws.cell(rr, 6, "Atención al cliente")
ws.cell(rr + 2, 1, "Costo por lead calificado (presupuesto año 1 / MQL)").font = Font(bold=True)
ws.cell(rr + 2, 2, "=Presupuesto!B13/B6").number_format = '"US$"#,##0'
ws.cell(rr + 3, 1, "Costo de marketing por línea vendida").font = Font(bold=True)
ws.cell(rr + 3, 2, "=Presupuesto!B13/B9").number_format = '"US$"#,##0'

# 13. Presupuesto con fórmulas
ws = sheet(wb, "Presupuesto", "Presupuesto de marketing (estimación, US$)", "Edite las cifras amarillas; totales y porcentajes se recalculan.")
items = [("Identidad, web, catálogo y fichas técnicas", 6000, 3000, 4000), ("Video y fotografía (efectos, demos, casos)", 5000, 6000, 6000),
         ("Kits de muestras ABM y muestrario", 6000, 7000, 8000), ("Ferias (stand, traslado de máquina, viáticos)", 18000, 25000, 32000),
         ("Inauguración sala demo y Días Híbridos", 6000, 8000, 8000), ("Publicidad digital (LinkedIn, Meta, YouTube, Google)", 5000, 9000, 12000),
         ("Herramientas (CRM, WhatsApp, correo)", 1500, 2000, 3000), ("Imprevistos", 2500, 5000, 7000)]
table(ws, 4, ["Concepto", "Año 1", "Año 2", "Año 3", "% año 1"], [[a, b, c, d, ""] for a, b, c, d in items], [52, 14, 14, 14, 12])
for i in range(len(items)):
    rr = 5 + i
    for cc in (2, 3, 4): ws.cell(rr, cc).number_format = '"US$"#,##0'; ws.cell(rr, cc).fill = INPUT
    ws.cell(rr, 5, f"=IFERROR(B{rr}/B$13,0)").number_format = "0%"
ws.cell(13, 1, "Total").font = Font(bold=True)
for cc in (2, 3, 4):
    c = ws.cell(13, cc, f"=SUM({L(cc)}5:{L(cc)}12)"); c.font = Font(bold=True); c.number_format = '"US$"#,##0'
    c.fill = PatternFill("solid", fgColor=PINK)
ws.cell(15, 1, "Ya incluido en la inversión inicial de la propuesta (Lanzamiento)"); ws.cell(15, 2, 20000).number_format = '"US$"#,##0'; ws.cell(15, 2).fill = INPUT
ws.cell(16, 1, "Fondo cooperativo CLT (3–5% de compras) — supuesto 4%"); ws.cell(16, 3, 0.04).number_format = "0%"; ws.cell(16, 3).fill = INPUT
ws.cell(17, 1, "Compras del distribuidor a CLT año 1 (capturar)"); ws.cell(17, 2, 0).number_format = '"US$"#,##0'; ws.cell(17, 2).fill = INPUT
ws.cell(18, 1, "Aportación estimada de CLT"); ws.cell(18, 2, "=B17*C16").number_format = '"US$"#,##0'
ws.cell(19, 1, "Gasto nuevo a cargo del grupo año 1").font = Font(bold=True)
ws.cell(19, 2, "=MAX(0,B13-B15-B18)").number_format = '"US$"#,##0'; ws.cell(19, 2).font = Font(bold=True)
ws.cell(21, 1, "Distribución mensual año 1 (US$)").font = Font(bold=True, size=12, color="C2185B")
month_w = [6, 8, 9, 7, 9, 10, 8, 12, 9, 7, 8, 7]  # peso relativo por mes (lanzamiento y FESPA)
for m in range(12):
    ws.cell(22, 2 + m, f"M{m + 1}").font = Font(bold=True)
    ws.cell(23, 2 + m, f"=$B$13*{month_w[m]}/{sum(month_w)}").number_format = '#,##0'
ws.cell(23, 1, "Gasto planeado")
ws.cell(24, 1, "Gasto real (capturar)")
for m in range(12): ws.cell(24, 2 + m).fill = INPUT; ws.cell(24, 2 + m).number_format = '#,##0'
for m in range(12): ws.column_dimensions[L(6 + m)].width = max(ws.column_dimensions[L(6 + m)].width or 0, 9)

# 14. KPIs mensuales (plan vs real)
ws = sheet(wb, "KPIs mensuales", "Tablero mensual de indicadores", "Meta mensual sugerida (año 1) y columna para capturar el real de cada mes.")
kpis = [("Contactos generados", 75), ("Leads calificados (MQL)", 25), ("Demos o pruebas", 3.3), ("Propuestas", 1), ("Líneas vendidas", 0.17),
        ("Talleres nuevos con consumibles", 2.5), ("Ventas de consumibles (US$)", None), ("Seguidores nuevos (Instagram + LinkedIn)", 150),
        ("Visitas a la página /clt", 600), ("Mensajes de WhatsApp recibidos", 60), ("Costo por lead (US$)", None), ("NPS clientes (90 días)", None)]
ws.cell(4, 1, "Indicador").font = HEAD; ws.cell(4, 1).fill = HFILL; ws.cell(4, 2, "Meta/mes").font = HEAD; ws.cell(4, 2).fill = HFILL
for m in range(12):
    c = ws.cell(4, 3 + m, f"Mes {m + 1}"); c.font = HEAD; c.fill = HFILL
c = ws.cell(4, 15, "Total año"); c.font = HEAD; c.fill = HFILL
c = ws.cell(4, 16, "Meta año"); c.font = HEAD; c.fill = HFILL
for i, (k, v) in enumerate(kpis):
    rr = 5 + i
    ws.cell(rr, 1, k).font = Font(bold=True); ws.cell(rr, 2, v if v is not None else "—")
    for m in range(12): ws.cell(rr, 3 + m).fill = INPUT; ws.cell(rr, 3 + m).border = BOX
    if v is not None:
        ws.cell(rr, 15, f"=SUM(C{rr}:N{rr})"); ws.cell(rr, 16, f"=ROUND(B{rr}*12,0)")
ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 11
for m in range(14): ws.column_dimensions[L(3 + m)].width = 9
ws.freeze_panes = "C5"

# 15. Calendario tipo Gantt (36 meses)
ws = sheet(wb, "Calendario", "Calendario de actividades (Gantt, 36 meses)", "Cambie inicio/fin y el gráfico se actualiza.")
acts = [("0. Preparar", "Kit de marca, web /clt, catálogo, fichas, videos", 1, 3), ("0. Preparar", "CRM, WhatsApp Business, calculadora de retorno", 1, 2),
        ("1. Lanzar", "Envío de muestras ABM a las 17 cuentas", 2, 6), ("1. Lanzar", "Inauguración de la sala demo (Querétaro)", 4, 4),
        ("1. Lanzar", "Campaña LinkedIn + Instagram", 4, 12), ("1. Lanzar", "Lanzamiento de consumibles a talleres", 4, 6),
        ("1. Lanzar", "FESPA México", 8, 8), ("1. Lanzar", "Programa «Clientes fundadores»", 5, 12),
        ("2. Probar y escalar", "Webinars mensuales / Academia Híbrida", 5, 36), ("2. Probar y escalar", "Días Híbridos: Guadalajara y León", 8, 9),
        ("2. Probar y escalar", "Días Híbridos: Puebla y Tijuana", 11, 12), ("2. Probar y escalar", "Días Híbridos: Monterrey y Mérida", 14, 15),
        ("2. Probar y escalar", "Casos de éxito en video", 10, 36), ("2. Probar y escalar", "Programa de referidos", 12, 36),
        ("2. Probar y escalar", "PRINTING United (con VSP)", 12, 12), ("3. Liderar", "Distribuidores de química (Guadalajara, Puebla, León)", 13, 24),
        ("3. Liderar", "FESPA México (año 2)", 20, 20), ("3. Liderar", "Intermoda (Guadalajara)", 18, 19),
        ("3. Liderar", "PRINTING United (año 2)", 24, 24), ("Latam", "Ola 1: El Salvador, Honduras, Guatemala", 20, 36),
        ("Latam", "Apparel Sourcing Show (Guatemala)", 22, 22), ("Latam", "Ola 2: Perú y Colombia", 25, 36),
        ("Latam", "Colombiatex (Medellín)", 26, 26), ("Latam", "Perú Moda (Lima)", 28, 28)]
hdr = ["Fase", "Actividad", "Inicio (mes)", "Fin (mes)"] + [str(m) for m in range(1, 37)]
for j, h in enumerate(hdr, 1):
    c = ws.cell(4, j, h); c.font = HEAD; c.fill = HFILL; c.alignment = Alignment(horizontal="center")
for i, (f, a, s, e) in enumerate(acts):
    rr = 5 + i
    ws.cell(rr, 1, f); ws.cell(rr, 2, a); ws.cell(rr, 3, s).fill = INPUT; ws.cell(rr, 4, e).fill = INPUT
    for m in range(36): ws.cell(rr, 5 + m).border = BOX
last = 4 + len(acts)
rng = f"E5:{L(4 + 36)}{last}"
ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(E$4*1>=$C5,E$4*1<=$D5)"], fill=PatternFill("solid", fgColor="C2185B")))
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 52; ws.column_dimensions["C"].width = 11; ws.column_dimensions["D"].width = 10
for m in range(36): ws.column_dimensions[L(5 + m)].width = 3.6
for j, (lab, s, e) in enumerate([("Año 1", 5, 16), ("Año 2", 17, 28), ("Año 3", 29, 40)]):
    ws.merge_cells(start_row=3, start_column=s, end_row=3, end_column=e)
    c = ws.cell(3, s, lab); c.font = Font(bold=True); c.alignment = Alignment(horizontal="center")
ws.freeze_panes = "E5"

# 16–18
for name, h1, title, widths in [
    ("Latinoamérica", "16.", "Expansión a Latinoamérica", [8, 12, 30, 50, 44]),
    ("Pedir a CLT", "17.", "Lo que pedimos a CLT para el marketing", [120]),
    ("Riesgos", "18.", "Riesgos del plan", [45, 70]),
]:
    ws = sheet(wb, name, title)
    blocks(ws, 3, section(h1), widths)

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = "C2185B" if ws.title in ("Resumen", "ABM 17 cuentas", "Presupuesto", "Calendario") else "1B1F3B"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

out = sys.argv[1] if len(sys.argv) > 1 else "Plan_Marketing_CLT_Mexico_Latam.xlsx"
wb.save(out)
print("ok", out, len(wb.worksheets), "hojas")
