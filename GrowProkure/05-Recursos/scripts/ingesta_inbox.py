#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ingesta R16: lee los Excel de _INBOX, normaliza al esquema R1, deduplica contra
todo lo que ya conoce el proyecto (R2), y corre la auditoria de riesgo de correos
con las mismas reglas que causaron el rebote de SIIX.

No envia nada. Escribe contactos nuevos + auditoria de riesgo.
"""
import openpyxl, csv, os, re, unicodedata
from collections import Counter, OrderedDict

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
INV, GTM = os.path.join(BASE,"02-Investigacion"), os.path.join(BASE,"04-GoToMarket")
INBOX = os.path.join(INV, "_INBOX")

# ---------- reglas de riesgo (heredadas de Estudios_Correos_En_Riesgo) ----------
GRATIS = {"gmail.com","outlook.com","outlook.es","hotmail.com","hotmail.es","hotmail.com.mx",
          "yahoo.com","yahoo.com.mx","yahoo.es","live.com","prodigy.net.mx","icloud.com","aol.com"}
GENERICO = {"info","ventas","compras","contacto","sales","admin","purchasing","procurement",
            "contact","hello","marketing","rh","atencion","purchase","buyer","office"}
EDU = re.compile(r"\.(edu|edu\.mx|ac\.[a-z]{2})$", re.I)
# TLD de matriz extranjera: si el contacto esta en MX/LATAM y el dominio es de otro pais => patron SIIX
EXTRANJERO = re.compile(r"\.(jp|kr|de|at|it|es|fr|se|ch|cn|tw|in|co\.jp|co\.kr|com\.tw|com\.cn)$", re.I)
LATAM = {"mexico","méxico","mx","brasil","brazil","colombia","chile","peru","perú","argentina",
         "costa rica","republica dominicana","república dominicana","guatemala","honduras","el salvador"}
RE_MAIL = re.compile(r"^[^@\s,;()<>]+@[^@\s,;()<>]+\.[a-z]{2,}$", re.I)

def norm(s):
    s = unicodedata.normalize("NFKD",(s or "")).encode("ascii","ignore").decode()
    return re.sub(r"[^a-z0-9]+"," ",s.lower()).strip()

def sheet(path, name):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if name not in wb.sheetnames: return []
    ws = wb[name]; it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    return [dict(zip(hdr, r)) for r in it]

def g(row, *names):
    for n in names:
        v = row.get(n)
        if v not in (None, ""): return str(v).strip()
    return ""

# ---------- lo que ya conocemos (R2) ----------
ya = set()
for f in ["Instantly_FILTRADO_COMPLETO.csv", "_Instantly_EXCLUIDOS_auditoria.csv"]:
    for r in csv.DictReader(open(os.path.join(GTM,f), encoding="utf-8")):
        if r["email"]: ya.add(r["email"].strip().lower())

FUENTES = [
 ("US_Electronics_SMT_Study_ACT.xlsx","Contacts","Electronica","US"),
 ("US_Electronics_SMT_Study_ACT.xlsx","Purchasing & Sourcing","Electronica","US"),
 ("Estudio_Electronica_Mexico_ESTRUCTURADO_ES.xlsx","Contactos","Electronica","MX"),
 ("Estudio_Electronica_Mexico_ESTRUCTURADO_ES.xlsx","Compras y Sourcing","Electronica","MX"),
 ("Estudio_Mercado_Electronica_LATAM_Caribe.xlsx","Contactos","Electronica","LATAM"),
]

nuevos, riesgo, dups = [], [], 0
vistos = set()
for arch, hoja, vertical, region in FUENTES:
    p = os.path.join(INBOX, arch)
    if not os.path.exists(p): continue
    for row in sheet(p, hoja):
        email = g(row,"Email","Correo","Email corporativo").lower()
        if not email or "@" not in email: continue
        if email in ya: dups += 1; continue
        if email in vistos: dups += 1; continue
        vistos.add(email)
        empresa = g(row,"Company","Empresa")
        pais    = g(row,"Pais","País") or ("Mexico" if region=="MX" else "United States" if region=="US" else "")
        ciudad  = g(row,"City","Ciudad","Ciudad / Hub")
        estado  = g(row,"State","Estado")
        nombre  = g(row,"Contact","Contacto","Nombre")
        puesto  = g(row,"Title","Puesto","Cargo")
        reg = OrderedDict(
            Vertical=vertical, Pais=pais, Lado="Demanda", Empresa=empresa,
            Contacto=nombre, Puesto=puesto,
            Funcion=g(row,"Target Function","Función Objetivo","Departamento","Role Category","Categoría de Rol"),
            Ciudad=ciudad, Estado=estado, Email=email,
            Telefono=g(row,"Direct Phone","Tel Directo","Telefono","Teléfono"),
            Segmento=g(row,"Vertical","Segmento"), Prioridad="",
            LinkedIn=g(row,"LinkedIn"), Fuente=arch, Origen=hoja,
            electronica=g(row,"Electronics?","¿Electrónica?"),
            nivel=g(row,"Mgmt Level","Nivel"), region=region)
        # --- auditoria de riesgo ---
        motivos = []
        if not RE_MAIL.match(email) or "inferido" in email: motivos.append("correo malformado o inferido")
        dom, user = email.split("@")[-1], email.split("@")[0]
        if dom in GRATIS: motivos.append("dominio personal/gratuito, no corporativo")
        if EDU.search(dom): motivos.append("dominio educativo, no es la empresa")
        if user in GENERICO: motivos.append("buzon generico, no nominal")
        ctx = norm(pais + " " + ciudad + " " + estado)
        if region in ("MX","LATAM") or any(k in ctx for k in ("mexico","brasil","colombia","chile")):
            if EXTRANJERO.search(dom): motivos.append("dominio de matriz extranjera en contacto LATAM (patron SIIX)")
        if len(norm(nombre).split()) <= 1 and norm(nombre) and norm(nombre) == norm(user):
            motivos.append("nombre = buzon (correo construido, no persona verificada)")
        if motivos:
            r2 = OrderedDict(reg); r2["motivo_riesgo"] = " | ".join(motivos)
            r2["severidad"] = "ALTA" if any(m.startswith(("correo malformado","dominio de matriz","dominio educativo")) for m in motivos) else "MEDIA"
            riesgo.append(r2)
        else:
            nuevos.append(reg)

def esc(p, rows, cols):
    with open(p,"w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=cols,extrasaction="ignore"); w.writeheader(); w.writerows(rows)

COLS=["Vertical","Pais","Lado","Empresa","Contacto","Puesto","Funcion","Ciudad","Estado",
      "Email","Telefono","Segmento","Prioridad","LinkedIn","Fuente","Origen",
      "electronica","nivel","region"]
esc(os.path.join(INV,"INBOX_Contactos_NUEVOS.csv"), nuevos, COLS)
esc(os.path.join(INV,"INBOX_Correos_En_Riesgo.csv"), riesgo, COLS+["severidad","motivo_riesgo"])

print("Ya conocidos en el proyecto :", len(ya))
print("Duplicados descartados (R2) :", dups)
print("NUEVOS limpios              :", len(nuevos))
print("NUEVOS con riesgo           :", len(riesgo))
for k,v in Counter(r["severidad"] for r in riesgo).most_common(): print("     ",k,v)
print()
print("motivos de riesgo:")
for k,v in Counter(m for r in riesgo for m in r["motivo_riesgo"].split(" | ")).most_common():
    print("   %-58s %4d" % (k,v))
print()
print("nuevos limpios por region:", Counter(r["region"] for r in nuevos).most_common())
print("nuevos limpios por fuente:", Counter(r["Fuente"] for r in nuevos).most_common())
