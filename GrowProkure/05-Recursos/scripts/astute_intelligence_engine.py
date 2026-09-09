import openpyxl, re
from collections import defaultdict, Counter
U="/root/.claude/uploads/79b088b9-ecb6-5156-8b46-8be359f25fa8/"
OUT="05-Recursos/Astute_Intelligence_Engine.xlsx"

def s(v): return str(v).strip() if v not in (None,"") else ""
def keyco(name):
    n=str(name or "").lower()
    n=re.sub(r'[^a-z0-9 ]',' ',n)
    n=re.sub(r'\b(s a de c v|sa de cv|s de rl de cv|s de rl|inc|corp|corporation|company|co|llc|ltd|limited|de mexico|mexicana|mexico|sci|international|group|grupo|holdings?|the)\b',' ',n)
    return re.sub(r'\s+',' ',n).strip()
def valid_email(e):
    e=str(e or "").strip().lower()
    if "@" not in e or " " in e: return ""
    if "." not in e.split("@")[-1]: return ""
    return e
def find(low,*keys):
    for i,h in enumerate(low):
        if any(k in h for k in keys): return i
    return None

# ---------- COMPANIES (directorios nivel empresa) ----------
companies={}   # key -> dict
def reg_company(name, pais, ciudad, estado, tipo, sector, evidencia, confianza, web, fuente):
    k=keyco(name)
    if not k: return
    c=companies.get(k)
    if not c:
        c=companies[k]={"key":k,"display":s(name),"aliases":set(),"pais":pais,"ciudad":ciudad,
            "estado":estado,"tipo":tipo,"sector":sector,"evidencia":evidencia,"confianza":confianza,
            "web":web,"fuentes":set(),"import_ev":False}
    c["aliases"].add(s(name))
    if len(s(name))>len(c["display"]): c["display"]=s(name)
    for f,v in [("pais",pais),("ciudad",ciudad),("estado",estado),("tipo",tipo),("sector",sector),
                ("evidencia",evidencia),("confianza",confianza),("web",web)]:
        if not c[f] and v: c[f]=v
    if fuente: c["fuentes"].add(fuente)

def load_dir(path, sheet, pais_fixed=None):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    if sheet not in wb.sheetnames: wb.close(); return
    ws=wb[sheet]; hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    low=[str(h).lower() if h else "" for h in hdr]
    ci={"emp":find(low,"empresa","company"),"pais":find(low,"pais","country"),
        "ciudad":find(low,"ciudad","city","hub"),"estado":find(low,"estado","state"),
        "tipo":find(low,"tipo","type"),"sector":find(low,"sector","segmento","vertical"),
        "ev":find(low,"evidencia","evidence"),"conf":find(low,"confianza","confidence"),
        "web":find(low,"sitio","website","web")}
    def g(r,k): i=ci[k]; return s(r[i]) if (i is not None and i<len(r)) else ""
    PAIS={"brasil":"BR","argentina":"AR","costa rica":"CR","rep. dominicana":"DO","colombia":"CO","puerto rico":"PR","chile":"CL","peru":"PE","ecuador":"EC"}
    for r in ws.iter_rows(min_row=2,values_only=True):
        if not g(r,"emp"): continue
        pais=pais_fixed or PAIS.get(g(r,"pais").lower(), g(r,"pais")[:2].upper())
        reg_company(g(r,"emp"),pais,g(r,"ciudad"),g(r,"estado"),g(r,"tipo"),g(r,"sector"),g(r,"ev"),g(r,"conf"),g(r,"web"),sheet)
    wb.close()

load_dir(U+"ac76eb9e-Estudio_Electronica_Mexico_ESTRUCTURADO_ES.xlsx","Directorio",pais_fixed="MX")
load_dir(U+"6007af24-US_Electronics_SMT_Study_ACT.xlsx","Directory",pais_fixed="US")
load_dir(U+"5cbc9797-Estudio_Mercado_Electronica_LATAM_Caribe.xlsx","Directorio Empresas")

# ---------- IMPORT EVIDENCE ----------
def load_import(path, sheet):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    if sheet not in wb.sheetnames: wb.close(); return
    ws=wb[sheet]; hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    ci=find([str(h).lower() if h else "" for h in hdr],"empresa","company")
    for r in ws.iter_rows(min_row=2,values_only=True):
        if ci is not None and ci<len(r) and r[ci]:
            k=keyco(r[ci])
            if k in companies: companies[k]["import_ev"]=True
    wb.close()
load_import(U+"ac76eb9e-Estudio_Electronica_Mexico_ESTRUCTURADO_ES.xlsx","Evidencia Importación")
load_import(U+"6007af24-US_Electronics_SMT_Study_ACT.xlsx","Import Evidence")

# ---------- CONTACTS ----------
PURCH_A=("chief procurement","cpo","procurement director","sourcing director","director de compras","gerente de compras","purchasing manager","commodity manager","strategic sourcing","global sourcing","director of supply chain","supply chain director","director de abastecimiento","head of procurement","head of sourcing")
PURCH_B=("purchas","sourc","procure","commodity","material","supply chain","buyer","compra","abasto","planner","planning")
ENGf=("engineer","ingenier","smt","npi","manufactur","process","calidad","quality","operations","operaciones")
EXECf=("director","chief","vp","president","gerente general","ceo","coo","owner","dueñ","plant manager","gerente de planta","general manager")
def priority(title):
    t=str(title or "").lower()
    if any(x in t for x in PURCH_A): return "A"
    if any(x in t for x in PURCH_B): return "B"
    if any(x in t for x in EXECf): return "B"
    if any(x in t for x in ENGf): return "C"
    return "D"
def func_of(title):
    t=str(title or "").lower()
    if any(x in t for x in PURCH_B): return "Compras/Sourcing"
    if any(x in t for x in EXECf): return "Ejecutivo"
    if any(x in t for x in ENGf): return "Ingeniería/Manuf"
    return "Otro"

contacts=[]   # dicts
comp_contacts=defaultdict(list)
def load_contacts(path, sheets, pais_fixed=None):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    for sheet in sheets:
        if sheet not in wb.sheetnames: continue
        ws=wb[sheet]; hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
        low=[str(h).lower() if h else "" for h in hdr]
        ci={"emp":find(low,"empresa","company"),"name":find(low,"contacto","contact","nombre"),
            "title":find(low,"puesto","title","cargo"),"email":find(low,"correo","email"),
            "phone":find(low,"tel directo","direct phone","tel","phone"),
            "city":find(low,"ciudad","city"),"state":find(low,"estado","state"),
            "pais":find(low,"pais","country")}
        def g(r,k): i=ci[k]; return s(r[i]) if (i is not None and i<len(r)) else ""
        PAIS={"brasil":"BR","argentina":"AR","costa rica":"CR","rep. dominicana":"DO","colombia":"CO","puerto rico":"PR"}
        for r in ws.iter_rows(min_row=2,values_only=True):
            emp=g(r,"emp"); name=g(r,"name")
            if not emp and not name: continue
            em=valid_email(g(r,"email"))
            k=keyco(emp)
            pais=pais_fixed or PAIS.get(g(r,"pais").lower(), g(r,"pais")[:2].upper() if g(r,"pais") else "")
            rec={"company_key":k,"empresa":emp,"contacto":name,"puesto":g(r,"title"),
                 "email":em,"pri":priority(g(r,"title")),"func":func_of(g(r,"title")),
                 "ciudad":g(r,"city"),"estado":g(r,"state"),"pais":pais,"fuente":sheet}
            contacts.append(rec); comp_contacts[k].append(rec)
    wb.close()
load_contacts(U+"ac76eb9e-Estudio_Electronica_Mexico_ESTRUCTURADO_ES.xlsx",["Contactos","Compras y Sourcing"],pais_fixed="MX")
load_contacts(U+"6007af24-US_Electronics_SMT_Study_ACT.xlsx",["Contacts","Purchasing & Sourcing"],pais_fixed="US")
load_contacts(U+"5cbc9797-Estudio_Mercado_Electronica_LATAM_Caribe.xlsx",["Contactos"])

# companies that only appear via contacts (no directory row) -> register light
for k,lst in comp_contacts.items():
    if k and k not in companies:
        r=lst[0]
        reg_company(r["empresa"],r["pais"],r["ciudad"],r["estado"],"","","","","", "solo-contactos")

# ---------- TSUNAMI ----------
tsu_keys={};
wb=openpyxl.load_workbook(U+"6fb49a48-Clientes_CRM_Tsunami.xlsx",read_only=True,data_only=True)
ws=wb["Clientes CRM Tsunami"]; hdr=[str(c.value).lower() if c.value else "" for c in next(ws.iter_rows(min_row=1,max_row=1))]
ni=find(hdr,"nombre"); ti=find(hdr,"tipo")
tsu_list=[]
for r in ws.iter_rows(min_row=2,values_only=True):
    nm=s(r[ni]) if ni is not None and ni<len(r) else ""
    if not nm: continue
    k=keyco(nm)
    if k: tsu_keys[k]=s(r[ti]) if ti is not None and ti<len(r) else ""; tsu_list.append(k)
wb.close()
tsu_set=set(tsu_keys)
STOP={"de","la","el","los","del","and","the","group","grupo","mexico","usa","inc","corp"}
def sig(tokens): return {t for t in tokens if len(t)>2 and t not in STOP}
tsu_tokens={tk:sig(set(tk.split())) for tk in tsu_set}
def tsunami_match(k):
    if k in tsu_set: return "EXISTING CUSTOMER","HIGH"
    kt=sig(set(k.split()))
    if not kt: return "PROSPECT",""
    for tk,tts in tsu_tokens.items():
        if not tts: continue
        inter=kt & tts
        if inter and (inter==kt or inter==tts):   # una contiene a la otra (por tokens)
            if len(inter)>=2 or (len(inter)==1 and len(next(iter(inter)))>=7):
                return "EXISTING (posible)","POSSIBLE"
    return "PROSPECT",""
# marcadores de fabricante de componentes (proveedor, no comprador ideal para distribuidor)
def is_component_maker(c):
    t=(c["sector"]+" "+c["tipo"]).lower()
    return ("semiconductor" in t and any(x in t for x in ("oem","fab","idm","foundry","semi"))) or "wafer" in t

# ---------- SMT EVIDENCE SCORE ----------
def smt_score(c):
    conf=c["confianza"].lower(); ev=c["evidencia"].lower()
    has_ev = bool(ev) and ev not in ("info no disponible","info not found","","n/a")
    if conf in ("alta","high") and has_ev: return 5 if "smt" in ev or "pcba" in ev or "line" in ev else 4
    if conf in ("alta","high"): return 4
    if conf in ("media","medium"): return 3
    if conf in ("baja","low"): return 2
    if has_ev: return 2
    return 1 if c["tipo"] else 0

# ---------- INDUSTRY ATTRACTIVENESS ----------
HIGH_IND=("aeroespac","aerospace","defen","automotri","automotive","medic","medical","industrial","energ","semiconduct","instrument","telecom")
def industry_attr(c):
    t=(c["sector"]+" "+c["tipo"]).lower()
    if any(x in t for x in ("aeroespac","aerospace","defen","medic","medical","semiconduct")): return 10
    if any(x in t for x in ("automotri","automotive","industrial","energ","instrument","telecom")): return 8
    if any(x in t for x in ("consumo","consumer","white","linea blanca","línea blanca","electrodom")): return 5
    return 6 if c["tipo"] else 3
def eol_potential(c):
    t=(c["sector"]+" "+c["tipo"]).lower()
    return 5 if any(x in t for x in ("aeroespac","aerospace","defen","medic","medical","industrial","instrument")) else 3

# ---------- SCORE ----------
rows=[]
for k,c in companies.items():
    lst=comp_contacts.get(k,[])
    n_contacts=len(lst)
    n_purch=sum(1 for x in lst if x["func"]=="Compras/Sourcing")
    best=None
    for pr in ["A","B","C","D"]:
        cand=[x for x in lst if x["pri"]==pr and x["email"]]
        if cand: best=cand[0]; break
    if not best and lst: best=lst[0]
    smt=smt_score(c)
    ind=industry_attr(c)
    mfg = 15 if (c["tipo"] or smt>=2) else 5
    comp_intensity = round(ind*1.5)              # 0-15 from industry
    proc_complex = 10 if n_contacts>=5 else (6 if n_contacts>=2 else 3)
    scpain = (10 if c["import_ev"] else 5) if ind>=8 else (6 if c["import_ev"] else 3)
    scale = 10 if n_contacts>=8 else (7 if n_contacts>=3 else (4 if n_contacts>=1 else 2))
    imp = 5 if c["import_ev"] else 0
    eol = eol_potential(c)
    cavail = 5 if (best and best.get("email")) else (2 if best else 0)
    score = round(mfg*0.0+  # keep mfg separate below
                  0)
    score = mfg + round(smt/5*15) + comp_intensity + proc_complex + scpain + scale + ind + imp + eol + cavail
    maker=is_component_maker(c)
    if maker: score=max(0, score-25)   # penalización: proveedor, no comprador ideal
    score = min(score,100)
    tier = "A+" if score>=90 else "A" if score>=80 else "B" if score>=70 else "C" if score>=60 else "Low"
    match,conf_match=tsunami_match(k)
    rows.append({**c,"n_contacts":n_contacts,"n_purch":n_purch,"smt":smt,"ind":ind,"maker":maker,
        "score":score,"tier":tier,"tsunami":match,"tsunami_conf":conf_match,
        "best_name":best["contacto"] if best else "","best_title":best["puesto"] if best else "",
        "best_email":best.get("email","") if best else "","best_pri":best["pri"] if best else ""})

rows.sort(key=lambda x:-x["score"])
print("Empresas totales (Master DB):",len(rows))
print("Contactos totales:",len(contacts))
from collections import Counter
print("Por tier:",dict(Counter(r["tier"] for r in rows)))
print("Por Tsunami:",dict(Counter(r["tsunami"] for r in rows)))
print("Por país:",dict(Counter(r["pais"] or "?" for r in rows).most_common(12)))

# ---------- WRITE EXCEL ----------
wb=openpyxl.Workbook()
def sheet(title):
    ws=wb.create_sheet(title); return ws
# Master DB
ws=wb.active; ws.title="Master Company DB"
cols=["MasterID","Empresa","Pais","Estado","Ciudad","Tipo","Sector","SMT_Score(0-5)","ImportEv","#Contactos","#Compras","Tsunami","TsunamiConf","CompMaker","OpportunityScore","Tier","PrimaryContact","PrimaryTitle","PrimaryEmail","PrimaryPri","Web","Fuentes"]
ws.append(cols)
for r in rows:
    ws.append([r["key"],r["display"],r["pais"],r["estado"],r["ciudad"],r["tipo"],r["sector"],r["smt"],
        "Sí" if r["import_ev"] else "",r["n_contacts"],r["n_purch"],r["tsunami"],r["tsunami_conf"],
        "Sí (revisar fit)" if r["maker"] else "",r["score"],r["tier"],r["best_name"],r["best_title"],r["best_email"],r["best_pri"],r["web"],", ".join(sorted(r["fuentes"]))])
ws.auto_filter.ref=f"A1:{openpyxl.utils.get_column_letter(len(cols))}{ws.max_row}"
# Top 100 prospects (exclude existing customers for pure new-biz; keep existing as separate)
prospects=[r for r in rows if r["tsunami"].startswith("PROSPECT")]
top100=prospects[:100]
ws2=sheet("Top 100 Prospectos"); ws2.append(["Rank","Empresa","Pais","Estado","Industria/Tipo","SMT","Score","Tier","PrimaryContact","Role","Email","Why"])
for i,r in enumerate(top100,1):
    why=f"{r['tipo'] or 'Mfg'} · {r['sector'] or 'electronica'} · {r['n_contacts']} contactos" + (" · importador" if r["import_ev"] else "")
    ws2.append([i,r["display"],r["pais"],r["estado"],(r["tipo"]+" / "+r["sector"]).strip(" /"),r["smt"],r["score"],r["tier"],r["best_name"],r["best_pri"],r["best_email"],why])
# Top 25 actionable (prospect + has A/B purchasing contact w/ email)
actionable=[r for r in prospects if r["best_email"] and r["best_pri"] in ("A","B") and not r["maker"]][:25]
ws3=sheet("Top 25 Atacar Ya"); ws3.append(["Rank","Empresa","Pais","Estado","Industria","Score","Contact","Role","Email","Oportunidad(hipotesis)","Next Action"])
for i,r in enumerate(actionable,1):
    opp=[]
    ind_t=(r["sector"]+" "+r["tipo"]).lower()
    if any(x in ind_t for x in ("aeroespac","aerospace","defen","medic")): opp.append("EOL/LTB, obsolescence, hard-to-find")
    if any(x in ind_t for x in ("automotri","automotive")): opp.append("shortage mitigation, cost savings, alt sourcing")
    if r["import_ev"]: opp.append("alternative sourcing (importa)")
    if not opp: opp.append("cost savings, shortage support")
    ws3.append([i,r["display"],r["pais"],r["estado"],(r["tipo"]+" / "+r["sector"]).strip(" /"),r["score"],
        r["best_name"],r["best_pri"],r["best_email"]," + ".join(opp)+" [INFERRED]","Email personalizado + LinkedIn"])
# Existing customers found in studies (cross-sell / new plant / new contact)
existing=[r for r in rows if r["tsunami"].startswith("EXISTING")][:200]
ws4=sheet("Tsunami - Cuentas Existentes"); ws4.append(["Empresa","Pais","Match","Conf","#Contactos nuevos","PrimaryContact","Role","Email","Nota"])
for r in existing:
    ws4.append([r["display"],r["pais"],r["tsunami"],r["tsunami_conf"],r["n_contacts"],r["best_name"],r["best_pri"],r["best_email"],"Cross-sell / new plant / new buyer"])
# Posibles Tsunami (revisar manualmente)
posibles=[r for r in rows if r["tsunami_conf"]=="POSSIBLE"]
ws6=sheet("Posibles Tsunami (revisar)"); ws6.append(["Empresa","Pais","Score","Tier","PrimaryContact","Email","Nota"])
for r in posibles:
    ws6.append([r["display"],r["pais"],r["score"],r["tier"],r["best_name"],r["best_email"],"Revisar si es el mismo cliente de Tsunami antes de excluir del ranking"])
# Component makers (revisar fit — proveedores)
makers=[r for r in rows if r["maker"]]
ws7=sheet("Fabricantes componentes"); ws7.append(["Empresa","Pais","Tipo","Score(post-penal)","Nota"])
for r in makers:
    ws7.append([r["display"],r["pais"],r["tipo"],r["score"],"Fabricante de componentes/semis: proveedor, no comprador ideal. Fit sólo MRO/excedentes."])
# Data quality / methodology
ws5=sheet("Metodologia y Calidad")
notes=[
 ["ASTUTE INTELLIGENCE ENGINE — build automatizado sobre datos reales de los estudios + CRM Tsunami",""],
 ["Fecha","2026-08-27"],
 ["Fuentes","Estudio MX ESTRUCTURADO, US SMT ACT, LATAM_Caribe, CRM Tsunami (4659 clientes)"],
 ["Master Company ID","nombre normalizado (minúsculas, sin sufijos legales/geográficos). Merge por clave normalizada; aliases conservados."],
 ["SMT_Score (0-5)","Derivado de Confianza + Evidencia del directorio. 5=Alta+evidencia SMT/PCBA; 3=Media; 2=Baja/evidencia; 1=tipo sin evidencia; 0=nada. NO verificado en planta."],
 ["OpportunityScore (0-100)","mfg + smt(15) + intensidad-componentes(industria) + complejidad-compras(#contactos) + supply-pain + escala + industria + import(5) + EOL(5) + contacto(5). Es un score de PRIORIZACIÓN, no un hecho."],
 ["Tsunami match","exacto (HIGH) o contención de nombre (POSSIBLE). Revisar POSSIBLE manualmente antes de excluir."],
 ["Oportunidad/hipótesis","INFERRED a partir de industria/import. NO son hechos verificados."],
 ["Emails","tomados de los estudios (ZoomInfo/Apollo). Calidad no re-verificada aquí; nominal preferido. NUNCA inventados."],
 ["PENDIENTE (Fase externa)","Sales triggers en vivo (nuevas plantas, expansiones), verificación de email, ImportYeti/Panjiva real, plant-level detallado — hacer selectivo en Top 25."],
 ["NO destruir datos","La Base Maestra original se conserva; esto es una capa de inteligencia derivada."],
]
notes.append(["CALIBRACIÓN v2 (2026-08-29)","(1) Fabricantes de semiconductores/fabs penalizados -25 (proveedores, no compradores ideales) y excluidos del Top 25. (2) Match Tsunami endurecido a tokens (contención + 2 tokens signif. o 1 token largo) para reducir falsos 'posibles'. (3) Autofiltro en todas las hojas."])
for n in notes: ws5.append(n)

# autofiltro en todas las hojas tabulares
for wsx in wb.worksheets:
    if wsx.max_row>1 and wsx.title!="Metodologia y Calidad":
        wsx.auto_filter.ref=f"A1:{openpyxl.utils.get_column_letter(wsx.max_column)}{wsx.max_row}"

wb.save(OUT)
print("Fabricantes de componentes marcados:",sum(1 for r in rows if r["maker"]))
print("Posibles Tsunami (revisar):",sum(1 for r in rows if r["tsunami_conf"]=='POSSIBLE'))
print("\nGuardado:",OUT)
print("Top 100 prospectos:",len(top100),"| Top 25 accionables:",len(actionable),"| Existentes Tsunami:",len(existing))
print("\n--- TOP 15 PROSPECTOS ---")
for i,r in enumerate(top100[:15],1):
    print(f"{i:2d}. [{r['score']}|{r['tier']}] {r['display'][:34]:34} {r['pais']} · {(r['tipo'] or r['sector'])[:20]:20} · {r['n_contacts']}c · {r['best_pri']} {r['best_email'][:28]}")
