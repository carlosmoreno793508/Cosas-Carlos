# Genera el Excel interno de prospectos a partir de src/prospectos.json
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

d = json.load(open("src/prospectos.json", encoding="utf-8"))
ZONA = ["Estado de México y CDMX", "Jalisco", "Puebla y Tlaxcala", "Guanajuato", "Nuevo León", "Yucatán", "Querétaro", "La Laguna", "Aguascalientes"]
def zkey(z):
    for i, n in enumerate(ZONA):
        if n.split()[0].lower() in z.lower(): return i
    return 99
INK, PINK = "1B1F3B", "FCE4EC"
head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor=INK)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook(); ws = wb.active; ws.title = "Prospectos"
cols = ["Prioridad", "Zona", "Empresa", "Ciudad", "Dirección publicada", "A qué se dedica", "Subsector textil", "Dirección verificada", "Tipo", "Por qué es prospecto", "Sitio web", "Contacto publicado", "Fuente", "Estatus", "Responsable", "Próximo paso", "Notas"]
widths = [10, 22, 28, 24, 45, 50, 32, 12, 24, 60, 32, 24, 32, 16, 16, 24, 30]
ws.append(cols)
rows = sorted(d["prospects"], key=lambda p: ({"A": 0, "B": 1, "C": 2}.get(p["priority"], 3), zkey(p["zone"]), p["company"]))
for p in rows:
    ws.append([p["priority"], p["zone"], p["company"], p["city"], p.get("address", ""), p.get("activity", ""), p.get("subsector", ""), {True: "Sí", False: "No"}.get(p.get("verified"), ""), p["type"], p["why"], p["website"], p["phone_or_email"], p["source"], "Por contactar", "", "", p.get("flag", "")])
for i, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
for c in ws[1]: c.font, c.fill = head, fill
for r in ws.iter_rows(min_row=2):
    for c in r: c.alignment = wrap
    if r[0].value == "A":
        for c in r[:13]: c.fill = PatternFill("solid", fgColor=PINK)
ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type="list", formula1='"Por contactar,Contactado,Visita agendada,Demo,Cotización,Ganado,Perdido"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"N2:N{ws.max_row}")

ws2 = wb.create_sheet("Resumen por zona")
ws2.append(["Zona", "Total", "A", "B", "C"])
zs = {}
for p in d["prospects"]:
    z = zs.setdefault(p["zone"], {"A": 0, "B": 0, "C": 0}); z[p["priority"]] = z.get(p["priority"], 0) + 1
for z in sorted(zs, key=zkey):
    v = zs[z]; ws2.append([z, sum(v.values()), v["A"], v["B"], v["C"]])
n = ws2.max_row
ws2.append(["Total", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})"])
for c in ws2[1]: c.font, c.fill = head, fill
for c in ws2[ws2.max_row]: c.font = Font(bold=True)
ws2.column_dimensions["A"].width = 30

ws3 = wb.create_sheet("Competidores")
ws3.append(["Empresa", "Ciudad", "Qué vende", "Fuente"])
for c in d["competitors_seen"]: ws3.append([c["company"], c["city"], c["what"], c["source"]])
for c in ws3[1]: c.font, c.fill = head, fill
for col, w in zip("ABCD", [28, 20, 60, 40]): ws3.column_dimensions[col].width = w
for r in ws3.iter_rows(min_row=2):
    for c in r: c.alignment = wrap

ws4 = wb.create_sheet("Notas")
ws4.append(["Notas y advertencias"]); ws4["A1"].font = Font(bold=True)
for t in d.get("notes", []) + [
    "Prioridad A: estampado en volumen confirmado o fuerte; B: buen perfil; C: posible.",
    "Solo ADN Serigrafía y KW Puebla confirman pulpos automáticos en su propio sitio; el resto debe validarse por llamada.",
    "Uso interno de TID / VSP / CEB. La hoja 80-20 sí aparece en la propuesta para socios (sección 7).",
]:
    ws4.append([t])
ws4.column_dimensions["A"].width = 120
for r in ws4.iter_rows(min_row=2):
    for c in r: c.alignment = Alignment(wrap_text=True)

ws5 = wb.create_sheet("OEM", 1)
oc = ["Prioridad", "Zona", "Empresa", "Ciudad", "Dirección publicada", "A qué se dedica", "Marcas que surte (según fuente)", "Subsector textil", "Dirección verificada", "Contacto publicado", "Fuente", "Notas"]
ws5.append(oc)
oem = sorted([p for p in d["prospects"] if p.get("segment") == "OEM"], key=lambda p: ({"A": 0, "B": 1, "C": 2}.get(p["priority"], 3), zkey(p["zone"]), p["company"]))
for p in oem:
    ws5.append([p["priority"], p["zone"], p["company"], p["city"], p.get("address", ""), p.get("activity", ""), p.get("brands_served", ""), p.get("subsector", ""), {True: "Sí", False: "No"}.get(p.get("verified"), ""), p.get("phone_or_email", ""), p.get("source", ""), p.get("flag", "")])
for c in ws5[1]: c.font, c.fill = head, fill
for col, w in zip("ABCDEFGHIJKL", [10, 24, 34, 22, 45, 50, 36, 32, 12, 22, 36, 30]): ws5.column_dimensions[col].width = w
for r in ws5.iter_rows(min_row=2):
    for c in r: c.alignment = wrap
    if r[0].value == "A":
        for c in r: c.fill = PatternFill("solid", fgColor=PINK)
ws5.freeze_panes = "D2"; ws5.auto_filter.ref = ws5.dimensions

# Hoja 80/20: las 17 cuentas prioritarias de la propuesta (src/top8020.js)
import subprocess
top = json.loads(subprocess.check_output(["node", "-e", "console.log(JSON.stringify(require('./src/top8020')))"]))
by_name = {p["company"]: p for p in d["prospects"]}
def find(name):
    key = name.split(" (")[0].lower()
    return next((p for n, p in by_name.items() if n.lower().startswith(key)), {})
ws8 = wb.create_sheet("80-20", 0)
ws8.append(["#", "Segmento", "Empresa", "Ciudad", "Zona", "A qué se dedica", "Marcas y clientes / subsector", "Dirección publicada", "Sitio web", "Estatus"])
for i, (kind, lst) in enumerate([("Prospecto", top["prospects"]), ("OEM", top["oems"])]):
    for r in lst:
        p = find(r[0])
        ws8.append([ws8.max_row, kind, r[0], r[1], p.get("zone", ""), r[2], r[4], p.get("address", ""), p.get("website", ""), "Por contactar"])
for c in ws8[1]: c.font, c.fill = head, fill
for col, w in zip("ABCDEFGHIJ", [5, 12, 34, 22, 26, 60, 34, 45, 32, 16]): ws8.column_dimensions[col].width = w
for r in ws8.iter_rows(min_row=2):
    for c in r: c.alignment = wrap
ws8.freeze_panes = "D2"; ws8.auto_filter.ref = ws8.dimensions
dv8 = DataValidation(type="list", formula1='"Por contactar,Contactado,Visita agendada,Demo,Cotización,Ganado,Perdido"', allow_blank=True)
ws8.add_data_validation(dv8); dv8.add(f"J2:J{ws8.max_row}")

import os
lat = []
for f in ["latam_centroamerica.json", "latam_caribe.json", "latam_sudamerica.json"]:
    fp = os.path.join("src", f)
    if os.path.exists(fp):
        lat += json.load(open(fp, encoding="utf-8"))["prospects"]
if lat:
    ws6 = wb.create_sheet("Latinoamérica", 2)
    lc = ["Prioridad", "Región", "País", "Empresa", "Ciudad", "Tipo", "Dirección publicada", "A qué se dedica", "Marcas que surte (según fuente)", "Subsector textil", "Segmento", "Verificado", "Sitio web", "Contacto publicado", "Por qué es prospecto", "Fuente", "Estatus"]
    ws6.append(lc)
    ro = {"Centroamérica": 0, "Caribe": 1, "Sudamérica": 2}
    lat.sort(key=lambda p: ({"A": 0, "B": 1, "C": 2}.get(p.get("priority"), 3), ro.get(p.get("region"), 9), p.get("country", ""), p.get("company", "")))
    for p in lat:
        ws6.append([p.get("priority"), p.get("region"), p.get("country"), p.get("company"), p.get("city"), p.get("type"), p.get("address"), p.get("activity"), p.get("brands_served"), p.get("subsector"), p.get("segment"), {True: "Sí", False: "No"}.get(p.get("verified"), ""), p.get("website"), p.get("phone_or_email"), p.get("why"), p.get("source"), "Por contactar"])
    for c in ws6[1]: c.font, c.fill = head, fill
    for col, w in zip("ABCDEFGHIJKLMNOPQ", [10, 16, 18, 34, 18, 26, 45, 50, 36, 30, 11, 11, 30, 24, 50, 36, 14]): ws6.column_dimensions[col].width = w
    for r in ws6.iter_rows(min_row=2):
        for c in r: c.alignment = wrap
        if r[0].value == "A":
            for c in r: c.fill = PatternFill("solid", fgColor=PINK)
    ws6.freeze_panes = "E2"; ws6.auto_filter.ref = ws6.dimensions
    dv2 = DataValidation(type="list", formula1='"Por contactar,Contactado,Visita agendada,Demo,Cotización,Ganado,Perdido"', allow_blank=True)
    ws6.add_data_validation(dv2); dv2.add(f"Q2:Q{ws6.max_row}")
    ws7 = wb.create_sheet("Resumen Latam", 3)
    ws7.append(["Región", "País", "A", "B", "C", "Total"])
    cnt = {}
    for p in lat:
        k = (p.get("region"), p.get("country")); cnt.setdefault(k, {"A": 0, "B": 0, "C": 0}); cnt[k][p.get("priority", "C")] = cnt[k].get(p.get("priority", "C"), 0) + 1
    for (rg, ct), v in sorted(cnt.items(), key=lambda x: (ro.get(x[0][0], 9), -sum(x[1].values()))):
        ws7.append([rg, ct, v["A"], v["B"], v["C"], v["A"] + v["B"] + v["C"]])
    n = ws7.max_row
    ws7.append(["Total", "", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})"])
    for c in ws7[1]: c.font, c.fill = head, fill
    for c in ws7[ws7.max_row]: c.font = Font(bold=True)
    ws7.column_dimensions["A"].width = 18; ws7.column_dimensions["B"].width = 22

wb.save(sys.argv[1]); print("ok", sys.argv[1], len(rows))
